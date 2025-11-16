import sys
import os
import tempfile
import json

# 确保当前目录在Python路径中
sys.path.append('.')

# 导入节点
from text_save_nodes import Qwen3_Text_Save

def test_node_import():
    """测试节点导入是否成功"""
    try:
        node = Qwen3_Text_Save()
        print("✓ 节点导入成功")
        return True
    except Exception as e:
        print(f"✗ 节点导入失败: {e}")
        return False

def test_json_formatting():
    """测试JSON字符串格式化功能"""
    try:
        node = Qwen3_Text_Save()
        
        # 测试JSON字符串
        test_json = '''{
            "key": "value",
            "number": 123,
            "list": [1, 2, 3]
        }'''
        
        formatted = node.format_json_to_line(test_json)
        print(f"✓ JSON格式化成功: {formatted}")
        
        # 验证格式化后的字符串是有效的JSON
        json.loads(formatted)
        print("✓ 格式化后的JSON是有效的")
        
        # 测试非JSON字符串
        test_text = "这是一段普通文本"
        formatted_text = node.format_json_to_line(test_text)
        print(f"✓ 非JSON文本处理成功: {formatted_text}")
        
        return True
    except Exception as e:
        print(f"✗ JSON格式化失败: {e}")
        return False

def test_txt_save():
    """测试TXT文件保存功能"""
    try:
        node = Qwen3_Text_Save()
        
        # 创建临时文件
        with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False) as f:
            temp_file_path = f.name
        
        # 测试写入模式
        test_text = "这是第一行文本"
        status, file_path = node.save_text(test_text, temp_file_path, "txt", "write")
        print(f"✓ TXT写入模式: {status}")
        
        # 验证文件内容
        with open(temp_file_path, 'r', encoding='utf-8') as f:
            content = f.read().strip()
        if content == test_text:
            print("✓ TXT写入内容正确")
        else:
            print(f"✗ TXT写入内容错误: 预期 '{test_text}', 实际 '{content}'")
            return False
        
        # 测试追加模式
        test_text2 = "这是第二行文本"
        status, file_path = node.save_text(test_text2, temp_file_path, "txt", "append")
        print(f"✓ TXT追加模式: {status}")
        
        # 验证追加内容
        with open(temp_file_path, 'r', encoding='utf-8') as f:
            lines = [line.strip() for line in f.readlines()]
        if lines == [test_text, test_text2]:
            print("✓ TXT追加内容正确")
        else:
            print(f"✗ TXT追加内容错误: {lines}")
            return False
        
        # 测试JSON写入TXT
        test_json = '{"name": "test", "value": 123}'
        status, file_path = node.save_text(test_json, temp_file_path, "txt", "append")
        print(f"✓ JSON写入TXT: {status}")
        
        # 清理临时文件
        os.unlink(temp_file_path)
        
        return True
    except Exception as e:
        print(f"✗ TXT保存失败: {e}")
        return False

def test_excel_save():
    """测试Excel文件保存功能"""
    try:
        # 检查openpyxl是否已安装
        import openpyxl
        from openpyxl import Workbook, load_workbook
        
        node = Qwen3_Text_Save()
        
        # 创建临时文件路径（不创建实际文件，由openpyxl处理）
        import tempfile
        temp_file_path = tempfile.mktemp(suffix='.xlsx')
        
        # 测试写入模式
        test_text = "Excel写入测试"
        status, file_path = node.save_text(test_text, temp_file_path, "excel", "write")
        print(f"✓ Excel写入模式: {status}")
        
        # 验证写入内容
        wb = load_workbook(temp_file_path)
        ws = wb.active
        cell_value = ws.cell(row=1, column=1).value
        if cell_value == test_text:
            print("✓ Excel写入内容正确")
        else:
            print(f"✗ Excel写入内容错误: 预期 '{test_text}', 实际 '{cell_value}'")
            return False
        
        # 测试追加模式
        test_text2 = "Excel追加测试"
        status, file_path = node.save_text(test_text2, temp_file_path, "excel", "append")
        print(f"✓ Excel追加模式: {status}")
        
        # 验证追加内容
        wb = load_workbook(temp_file_path)
        ws = wb.active
        cell_value1 = ws.cell(row=1, column=1).value
        cell_value2 = ws.cell(row=2, column=1).value
        if cell_value1 == test_text and cell_value2 == test_text2:
            print("✓ Excel追加内容正确")
        else:
            print(f"✗ Excel追加内容错误: 行1='{cell_value1}', 行2='{cell_value2}'")
            return False
        
        # 测试JSON写入Excel
        test_json = '{"excel": "json test", "value": 456}'
        status, file_path = node.save_text(test_json, temp_file_path, "excel", "append")
        print(f"✓ JSON写入Excel: {status}")
        
        # 清理临时文件
        os.unlink(temp_file_path)
        
        return True
    except ImportError:
        print("⚠ Excel测试跳过: openpyxl未安装")
        return True  # 安装依赖不是测试重点
    except Exception as e:
        print(f"✗ Excel保存失败: {e}")
        return False

def main():
    """主测试函数"""
    print("开始测试Qwen3文本保存节点...")
    print("=" * 50)
    
    tests = [
        test_node_import,
        test_json_formatting,
        test_txt_save,
        test_excel_save
    ]
    
    passed = 0
    failed = 0
    
    for test in tests:
        if test():
            passed += 1
        else:
            failed += 1
        print("-" * 50)
    
    print(f"测试完成: {passed} 项通过, {failed} 项失败")
    
    if failed == 0:
        print("🎉 所有测试通过! 节点功能正常")
        return 0
    else:
        print("❌ 部分测试失败")
        return 1

if __name__ == "__main__":
    sys.exit(main())